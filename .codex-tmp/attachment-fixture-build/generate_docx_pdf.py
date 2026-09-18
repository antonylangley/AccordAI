import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


OUTPUT_DIR = Path("/Users/antonylangley/Projects/accord/test-fixtures/attachment-uploads")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

NAVY = "0B172B"
PURPLE = "635BFF"
MUTED = "667085"
LIGHT = "EEF1F6"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def style_run(run, size=10, bold=False, color=NAVY) -> None:
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def create_docx() -> None:
    path = OUTPUT_DIR / "05-redact-project-contacts.docx"
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    style_run(title.add_run("Project Atlas contacts"), size=18, bold=True)

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(12)
    style_run(
        intro.add_run(
            "Synthetic test data for validating local attachment redaction. "
            "Summarize the contacts and their assigned workstreams."
        ),
        size=10,
        color=MUTED,
    )

    rows = [
        ["Name", "Email", "Phone", "Workstream"],
        ["Jordan Rivera", "jordan.rivera@example.test", "202-555-0147", "Vendor onboarding"],
        ["Maya Chen", "maya.chen@example.test", "202-555-0182", "Release readiness"],
        ["Noah Williams", "noah.williams@example.test", "202-555-0196", "Documentation"],
    ]
    table = doc.add_table(rows=len(rows), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = [Inches(1.35), Inches(2.15), Inches(1.35), Inches(1.75)]
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.width = widths[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            style_run(
                paragraph.add_run(value),
                size=9,
                bold=row_index == 0,
                color="FFFFFF" if row_index == 0 else NAVY,
            )
            if row_index == 0:
                set_cell_shading(cell, PURPLE)

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(12)
    style_run(
        note.add_run("Expected result: ALLOW after redaction. The uploaded governed copy should not contain the original names, emails, or phone numbers."),
        size=9,
        color=MUTED,
    )
    doc.save(path)


def create_pdf() -> None:
    path = OUTPUT_DIR / "07-block-unpublished-forecast.pdf"
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "FixtureTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=HexColor(f"#{NAVY}"),
        alignment=TA_LEFT,
        spaceAfter=12,
    )
    body_style = ParagraphStyle(
        "FixtureBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        leading=15,
        textColor=HexColor(f"#{NAVY}"),
        spaceAfter=10,
    )
    muted_style = ParagraphStyle(
        "FixtureMuted",
        parent=body_style,
        fontSize=9,
        textColor=HexColor(f"#{MUTED}"),
    )
    doc = SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
        title="Synthetic unpublished forecast attachment test",
        author="Accord attachment test fixture",
    )
    story = [
        Paragraph("Unpublished Q3 revenue forecast", title_style),
        Paragraph(
            "Restricted internal draft. Do not share with external AI services or anyone outside the finance leadership group.",
            body_style,
        ),
        Spacer(1, 4),
    ]
    data = [
        ["Business unit", "Forecast revenue", "Allocated spending", "Planning note"],
        ["Enterprise", "$18.4M", "$6.2M", "Delay two contractor starts"],
        ["Commercial", "$11.7M", "$4.8M", "Reduce event spending by $350K"],
        ["Platform", "$7.9M", "$5.1M", "Hold infrastructure reserve"],
    ]
    table = Table(data, colWidths=[1.35 * inch, 1.3 * inch, 1.35 * inch, 2.25 * inch], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), HexColor(f"#{PURPLE}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), HexColor("#FFFFFF")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, HexColor("#D8DEE9")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [HexColor("#FFFFFF"), HexColor("#F7F8FB")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    story.extend(
        [
            table,
            Spacer(1, 14),
            Paragraph(
                "Test prompt: Summarize the forecast and recommend where we should cut spending.",
                body_style,
            ),
            Paragraph(
                "Expected result: BLOCK. This file contains explicit unpublished financial forecasts and allocated spending.",
                muted_style,
            ),
        ]
    )
    doc.build(story)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("format", choices=("docx", "pdf"))
    args = parser.parse_args()
    if args.format == "docx":
        create_docx()
    else:
        create_pdf()
